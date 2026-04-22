#!/usr/bin/env python3
# XLSX/CSV -> JSON donusturucu (El Recodo tango veritabani icin)
# Kullanim:
#   python scripts/xlsx_to_json.py                         # varsayilan: el_recodo_Tum Kayitlar.xlsx
#   python scripts/xlsx_to_json.py path/to/file.xlsx
#   python scripts/xlsx_to_json.py path/to/file.csv -o out.json

import argparse
import csv
import json
import re
import sys
import unicodedata
from pathlib import Path

# Turkce XLSX kolonu -> JSON key eslemesi
COLUMN_MAP = {
    "Kayıt No": "id",
    "Tarih": "date",
    "Yıl": "year",
    "Başlık": "title",
    "Tarz": "genre",
    "Orkestra": "orchestra",
    "Şarkıcı": "singer",
    "Besteci": "composer",
    "Yazar": "lyricist",
    "Etiket": "label",
    "Süre": "duration",
    "Duygular/Etiketler": "tags",
    "Dinlemek /10": "listen_rating",
    "Dans /10": "dance_rating",
}

# Normalizasyon (content.js'deki ile ayni algoritma olmali)
_PUNCT_RE = re.compile(r"[^\w\s]", re.UNICODE)
_WS_RE = re.compile(r"\s+")

def normalize(text):
    """Aksan kaldir, kucuk harf yap, noktalama temizle."""
    if text is None:
        return ""
    s = str(text)
    # NFD bolumu + combining mark'leri at
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.lower()
    s = _PUNCT_RE.sub(" ", s)
    s = _WS_RE.sub(" ", s).strip()
    return s

def format_duration(val):
    """datetime.time veya string alip mm:ss formatina cevir."""
    if val is None or val == "":
        return ""
    # openpyxl time objesi
    if hasattr(val, "minute") and hasattr(val, "second"):
        total_seconds = val.hour * 3600 + val.minute * 60 + val.second
        m, s = divmod(total_seconds, 60)
        return f"{m}:{s:02d}"
    return str(val)

def to_int(val):
    try:
        return int(val) if val not in (None, "") else None
    except (ValueError, TypeError):
        return None

def to_float(val):
    try:
        return float(val) if val not in (None, "") else None
    except (ValueError, TypeError):
        return None

def build_record(row):
    """Bir satir dict'inden JSON kaydi olustur."""
    title = (row.get("title") or "").strip()
    orchestra = (row.get("orchestra") or "").strip()
    singer = (row.get("singer") or "").strip() or "Instrumental"
    genre = (row.get("genre") or "").strip().lower()

    record = {
        "id": (row.get("id") or "").strip() or None,
        "title": title,
        "orchestra": orchestra,
        "singer": singer,
        "composer": (row.get("composer") or "").strip(),
        "lyricist": (row.get("lyricist") or "").strip(),
        "year": to_int(row.get("year")),
        "date": (row.get("date") or "").strip() if isinstance(row.get("date"), str) else (str(row.get("date")) if row.get("date") is not None else ""),
        "genre": genre,
        "label": (row.get("label") or "").strip(),
        "duration": format_duration(row.get("duration")),
        "tags": (row.get("tags") or "").strip(),
        "listen_rating": to_float(row.get("listen_rating")),
        "dance_rating": to_float(row.get("dance_rating")),
        # Fuse.js runtime araması için önceden normalize edilmiş alanlar
        "_n_title": normalize(title),
        "_n_orchestra": normalize(orchestra),
        "_n_singer": normalize(singer if singer != "Instrumental" else ""),
    }
    # Birlesik arama alani - title + orkestra + sarkici (Fuse extended search icin)
    record["_search"] = (
        record["_n_title"] + " " + record["_n_orchestra"] + " " + record["_n_singer"]
    ).strip()
    return record

def read_xlsx(path):
    """XLSX dosyasini satir satir oku, eslenen anahtarlarla dict uret."""
    try:
        import openpyxl
    except ImportError:
        print("HATA: openpyxl kurulu degil. Yukle: pip install openpyxl", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(c) if c is not None else "" for c in next(rows_iter)]
    # Her xlsx kolon indexi -> JSON key
    index_to_key = {i: COLUMN_MAP[h] for i, h in enumerate(header) if h in COLUMN_MAP}
    for row in rows_iter:
        if row is None or all(c is None or c == "" for c in row):
            continue
        mapped = {}
        for i, key in index_to_key.items():
            if i < len(row):
                mapped[key] = row[i]
        yield mapped

def read_csv(path):
    """CSV dosyasini satir satir oku. Header ayni Turkce kolon adlari olmali."""
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            mapped = {}
            for h, v in row.items():
                if h in COLUMN_MAP:
                    mapped[COLUMN_MAP[h]] = v
            yield mapped

def main():
    parser = argparse.ArgumentParser(description="El Recodo XLSX/CSV -> JSON donusturucu")
    parser.add_argument("input", nargs="?", default="el_recodo_Tüm Kayıtlar.xlsx",
                        help="Girdi dosyasi (.xlsx veya .csv)")
    parser.add_argument("-o", "--output", default="tango-overlay/data/tango_database.json",
                        help="Cikti JSON yolu")
    args = parser.parse_args()

    inp = Path(args.input)
    if not inp.exists():
        print(f"HATA: Girdi dosyasi bulunamadi: {inp}", file=sys.stderr)
        sys.exit(1)

    ext = inp.suffix.lower()
    if ext == ".xlsx":
        rows = read_xlsx(inp)
    elif ext == ".csv":
        rows = read_csv(inp)
    else:
        print(f"HATA: Desteklenmeyen uzanti: {ext}. .xlsx veya .csv olmali.", file=sys.stderr)
        sys.exit(1)

    records = []
    for row in rows:
        rec = build_record(row)
        # Baslik bos olan kayitlari atla
        if not rec["title"]:
            continue
        records.append(rec)

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    # Kompakt yaz (eklenti paketi kucuk kalsin)
    with open(out, "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, separators=(",", ":"))

    size_kb = out.stat().st_size / 1024
    print(f"OK: {len(records)} kayit yazildi -> {out} ({size_kb:.1f} KB)")

if __name__ == "__main__":
    main()
